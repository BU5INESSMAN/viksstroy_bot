import asyncio
from io import BytesIO
from unittest.mock import AsyncMock, patch

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from test_smr_merged_objects_report import _database
from test_smr_ready_management import _make_db
from smr_data import get_smr_read_model
from web.routers import kp
from services.smr_layouts import SMR_LAYOUTS, generate_smr_layout_bytes, smr_layout_file_entries
from services.smr_layouts import _work_rows


def _values(wb):
    return [c.value for ws in wb for row in ws for c in row if c.value is not None]


@pytest.mark.parametrize("variant", SMR_LAYOUTS)
def test_layouts_are_read_only_preserve_totals_and_both_objects(variant):
    async def scenario():
        db = _database()
        changes = db.conn.raw.total_changes
        report = await get_smr_read_model(db, 235)
        blob, filename = await generate_smr_layout_bytes(db, 235, variant, include_financial=True, include_participant_salary=True)
        wb = load_workbook(BytesIO(blob), data_only=True)
        values = _values(wb)
        assert report['totals']['price'] in values
        assert report['totals']['hours'] in values
        assert 'Рабочий первого объекта' in values
        assert 'Рабочий второго объекта' in values
        assert any('Ливневая канализация · З-100826-01' == v for v in values)
        assert any('Водопровод · З-100826-02' == v for v in values)
        assert 'Шурфление' in values and 'Протаскивание трубы' in values
        assert 'З-100826-01' in filename and '10.08.2026' in filename
        assert db.conn.raw.total_changes == changes
        db.conn.raw.close()
    asyncio.run(scenario())


def test_unassigned_work_is_only_on_one_sheet_never_guessed_or_duplicated():
    async def scenario():
        db = _database()
        db.conn.raw.executescript("""
            INSERT INTO teams VALUES (6,'Бригада 6');
            UPDATE application_hours SET team_id=6 WHERE id=2;
            UPDATE application_kp SET team_id=6 WHERE id=2;
            INSERT INTO application_extra_works VALUES
              (9,235,NULL,NULL,'Совместная работа','м',7,10,20,NULL,0,'2026-08-10',100);
        """)
        blob, _ = await generate_smr_layout_bytes(db, 235, 'brigades', include_financial=True)
        wb = load_workbook(BytesIO(blob), data_only=True)
        assert _values(wb).count('Совместная работа') == 1
        assert 'Без распределения' in wb.sheetnames
        assert 'Совместная работа' in [c.value for row in wb['Без распределения'] for c in row]
        assert _values(wb).count('Рабочий первого объекта') == 1
        assert _values(wb).count('Рабочий второго объекта') == 1
        db.conn.raw.close()
    asyncio.run(scenario())


@pytest.mark.parametrize('variant', SMR_LAYOUTS)
def test_non_financial_exports_do_not_leak_catalog_or_participant_pay(variant):
    async def scenario():
        db = _database()
        db.conn.raw.execute('UPDATE application_hours SET participant_salary=123456')
        blob, _ = await generate_smr_layout_bytes(db,235,variant)
        values = _values(load_workbook(BytesIO(blob),data_only=True))
        assert 123456 not in values
        assert 'Цена за ед., ₽' not in values
        assert 'ЗП участника, ₽' not in values
        assert 'ЗП по расценкам работ' not in values
        db.conn.raw.close()
    asyncio.run(scenario())


def test_zero_hours_and_deleted_name_are_visible_without_changing_classic_read_model():
    async def scenario():
        db=_database()
        db.conn.raw.executescript("UPDATE application_hours SET hours=0 WHERE id=1; DELETE FROM team_members WHERE id=17;")
        report=await get_smr_read_model(db,235)
        assert len(report['hours'])==1
        blob,_=await generate_smr_layout_bytes(db,235,'roster')
        values=_values(load_workbook(BytesIO(blob),data_only=True))
        assert 'Имя не найдено (карточка 17)' in values
        assert 'Рабочий второго объекта' in values
        db.conn.raw.close()
    asyncio.run(scenario())


def test_service_name_cannot_become_excel_formula():
    async def scenario():
        db=_database()
        db.conn.raw.execute("UPDATE kp_catalog SET name='=1+1' WHERE id=145943")
        blob,_=await generate_smr_layout_bytes(db,235,'overview')
        wb=load_workbook(BytesIO(blob),data_only=False)
        matching=[c for ws in wb for row in ws for c in row if c.value=='=1+1']
        assert len(matching)==1 and matching[0].data_type=='s'
        db.conn.raw.close()
    asyncio.run(scenario())


def test_entries_have_three_unique_scopes_and_filenames():
    entries=smr_layout_file_entries(235,['Объект'])
    assert len(entries)==3
    assert len({e['download_url'] for e in entries})==3
    assert all(e['kind']=='comparison' for e in entries)


@pytest.mark.parametrize('user,allowed', [
    ({'role':'hr','tg_id':500},True),
    ({'role':'foreman','tg_id':100},True),
    ({'role':'foreman','tg_id':999},False),
    ({'role':'brigadier','tg_id':100},False),
    ({'role':'worker','tg_id':100},False),
])
def test_comparison_api_checks_ownership_and_roles(user,allowed):
    async def scenario():
        db=await _make_db()
        generator=AsyncMock(return_value=(b'xlsx','Отчёт.xlsx'))
        try:
            with patch.object(kp,'db',db), patch('services.smr_layouts.generate_smr_layout_bytes',generator):
                if allowed:
                    response=await kp.download_smr_report(1,team_id=None,scope='overview',current_user=user)
                    assert response.status_code==200
                    assert generator.await_args.kwargs['include_financial']==(user['role']=='hr')
                    assert generator.await_args.kwargs['include_participant_salary']
                else:
                    with pytest.raises(HTTPException) as error:
                        await kp.download_smr_report(1,team_id=None,scope='overview',current_user=user)
                    assert error.value.status_code==403
                    generator.assert_not_awaited()
        finally:
            await db.conn.close()
    asyncio.run(scenario())


def test_unknown_layout_is_rejected():
    async def scenario():
        with pytest.raises(ValueError):
            await generate_smr_layout_bytes(None,235,'invalid')
    asyncio.run(scenario())


def test_kopeck_rounding_preserves_canonical_total():
    report={'plan_works':[], 'extra_works':[
        {'volume':0.005,'salary':1,'price':1},
        {'volume':0.005,'salary':1,'price':1},
    ]}
    rows=_work_rows(report)
    assert sum(r['amount_price'] for r in rows)==0.01
    assert sum(r['amount_salary'] for r in rows)==0.01
    assert rows[1]['rounding_adjusted']

import unittest
from datetime import date
from io import BytesIO

import aiosqlite
from openpyxl import load_workbook

from web.services.smr_period_report import (
    _canonicalize_report_people,
    generate_period_report,
    load_period_data,
)


class SmrPeriodReportTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.conn = await aiosqlite.connect(":memory:")
        self.conn.row_factory = aiosqlite.Row
        await self.conn.executescript(
            """
            CREATE TABLE applications (
                id INTEGER PRIMARY KEY, date_target TEXT, smr_status TEXT,
                kp_status TEXT, object_id INTEGER, object_address TEXT,
                public_number TEXT, foreman_id INTEGER, foreman_name TEXT,
                smr_accounted_at TEXT, kp_archived INTEGER DEFAULT 0
            );
            CREATE TABLE objects (id INTEGER PRIMARY KEY, name TEXT);
            CREATE TABLE teams (id INTEGER PRIMARY KEY, name TEXT);
            CREATE TABLE users (user_id INTEGER PRIMARY KEY, fio TEXT);
            CREATE TABLE team_members (
                id INTEGER PRIMARY KEY, team_id INTEGER, fio TEXT, position TEXT
            );
            CREATE TABLE application_hours (
                id INTEGER PRIMARY KEY, app_id INTEGER, team_id INTEGER,
                user_id INTEGER, hours REAL, participant_salary REAL,
                is_additional INTEGER
            );
            CREATE TABLE kp_catalog (
                id INTEGER PRIMARY KEY, category TEXT, name TEXT, unit TEXT,
                salary REAL, price REAL
            );
            CREATE TABLE extra_works_catalog (
                id INTEGER PRIMARY KEY, name TEXT, unit TEXT, salary REAL, price REAL
            );
            CREATE TABLE application_kp (
                id INTEGER PRIMARY KEY, application_id INTEGER, kp_id INTEGER,
                volume REAL, unit TEXT, current_salary REAL, current_price REAL,
                is_additional INTEGER, team_id INTEGER
            );
            CREATE TABLE application_extra_works (
                id INTEGER PRIMARY KEY, application_id INTEGER, extra_work_id INTEGER,
                kp_id INTEGER, custom_name TEXT, unit TEXT, volume REAL,
                salary REAL, price REAL, is_additional INTEGER, team_id INTEGER
            );
            CREATE TABLE employee_identities (
                id INTEGER PRIMARY KEY, canonical_fio TEXT, normalized_fio TEXT,
                position TEXT, linked_user_id INTEGER, status TEXT
            );
            CREATE TABLE employee_identity_members (
                member_id INTEGER PRIMARY KEY, identity_id INTEGER,
                source_fio TEXT, source_position TEXT
            );

            INSERT INTO objects VALUES (1,'ВОС 1'),(2,'База');
            INSERT INTO teams VALUES (1,'Бригада А'),(2,'Бригада Б');
            INSERT INTO users VALUES (100,'Прораб Роман');
            INSERT INTO applications VALUES
                (1,'2026-08-10','approved',NULL,1,'','З-100826-01',100,'','2026-08-11 10:00',0),
                (2,'2026-08-11','approved',NULL,2,'','З-110826-01',100,'','2026-08-12 10:00',1),
                (3,'2026-08-12','approved',NULL,1,'','З-120826-01',100,'',NULL,0),
                (4,'2026-08-13','in_progress','submitted',1,'','З-130826-01',100,'','2026-08-14 10:00',0);
            INSERT INTO team_members VALUES
                (10,1,'Иванов Иван','Монтажник'),
                (11,1,'Петров Пётр','Сварщик'),
                (12,2,'Иванов Иван Иванович','Монтажник');
            INSERT INTO employee_identities VALUES
                (1,'Иванов Иван Иванович','иванов иван иванович','Монтажник',NULL,'active'),
                (2,'Петров Пётр','петров пётр','Сварщик',NULL,'active');
            INSERT INTO employee_identity_members VALUES
                (10,1,'Иванов Иван','Монтажник'),
                (12,1,'Иванов Иван Иванович','Монтажник'),
                (11,2,'Петров Пётр','Сварщик');
            INSERT INTO application_hours VALUES
                (1,1,1,10,8,4000,0),
                (2,1,1,11,4,0,0),
                (3,2,2,12,4,2000,0),
                (4,3,1,11,5,500,0),
                (5,4,1,10,99,9999,0);
            INSERT INTO kp_catalog VALUES (1,'Сети','Монтаж трубы','м',100,150);
            INSERT INTO extra_works_catalog VALUES (1,'Откачка','ч',50,80);
            INSERT INTO application_kp VALUES
                (1,1,1,10,'м',100,150,0,1),
                (2,2,1,2,'м',200,300,0,2),
                (3,3,1,1,'м',100,150,0,1);
            INSERT INTO application_extra_works VALUES
                (1,2,1,NULL,'','ч',1,50,80,0,2);
            """
        )
        self.db = type("Db", (), {"conn": self.conn})()

    async def asyncTearDown(self):
        await self.conn.close()

    async def test_default_includes_accounted_archived_and_deduplicates_identity(self):
        blob, filename, summary = await generate_period_report(
            self.db, date(2026, 8, 1), date(2026, 8, 31)
        )
        workbook = load_workbook(BytesIO(blob), data_only=True)
        self.assertEqual(workbook.sheetnames, ["Сводка", "По заявкам", "Работы"])
        self.assertEqual(workbook["Сводка"]["A1"].value, "Сводный отчёт СМР")
        self.assertIn("Архивные СМР включены", workbook["Сводка"]["A2"].value)
        names = [workbook["Сводка"].cell(row, 1).value for row in range(6, 8)]
        self.assertEqual(names.count("Иванов Иван Иванович"), 1)
        headers = [workbook["Сводка"].cell(5, column).value for column in range(1, workbook["Сводка"].max_column + 1)]
        self.assertNotIn("ЗП по работам", headers)
        self.assertIn("ЗП прораба", headers)
        detail_headers = [workbook["По заявкам"].cell(4, column).value for column in range(1, workbook["По заявкам"].max_column + 1)]
        self.assertNotIn("ЗП по работам", detail_headers)
        detail_numbers = {
            workbook["По заявкам"].cell(row, 2).value
            for row in range(5, workbook["По заявкам"].max_row + 1)
        }
        self.assertIn("З-110826-01", detail_numbers)  # archived
        self.assertNotIn("З-120826-01", detail_numbers)  # unaccounted
        self.assertEqual(summary["employees"], 2)
        self.assertEqual(summary["objects"], 2)
        self.assertEqual(summary["works"], 3)
        self.assertTrue(filename.endswith(".xlsx"))

    async def test_include_unaccounted_checkbox_adds_ready_reports_only(self):
        data = await load_period_data(
            self.db,
            date(2026, 8, 1),
            date(2026, 8, 31),
            include_unaccounted=True,
        )
        numbers = {row["public_number"] for row in data["hours"]}
        self.assertIn("З-120826-01", numbers)
        self.assertNotIn("З-130826-01", numbers)
        self.assertEqual(len(data["works"]), 4)

    def test_report_merges_only_unambiguous_short_names(self):
        rows = [
            {"member_id": 1, "member_fio": "Борисов Илья", "member_position": "Рабочий"},
            {"member_id": 2, "member_fio": "Борисов Илья Артемович", "member_position": "Монтажник"},
            {"member_id": 3, "member_fio": "Борисов Илья Артемович", "member_position": "Монтажник"},
            {"member_id": 4, "member_fio": "Иванов Иван", "member_position": ""},
            {"member_id": 5, "member_fio": "Иванов Иван Петрович", "member_position": ""},
            {"member_id": 6, "member_fio": "Иванов Иван Сергеевич", "member_position": ""},
        ]
        _canonicalize_report_people(rows)
        self.assertEqual(len({row["report_person_key"] for row in rows[:3]}), 1)
        self.assertTrue(all(row["member_fio"] == "Борисов Илья Артемович" for row in rows[:3]))
        self.assertNotEqual(rows[3]["report_person_key"], rows[4]["report_person_key"])
        self.assertNotEqual(rows[3]["report_person_key"], rows[5]["report_person_key"])


if __name__ == "__main__":
    unittest.main()

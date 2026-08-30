"""Read-only alternative SMR layouts. Never infer allocations or move staff.

Unlike the classic files, each object section puts people and all performed
works together. Shared legacy rows have no saved percentage allocation: they
are shown once, separately, instead of being silently divided or duplicated.
"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from decimal import Decimal, ROUND_HALF_UP
from smr_calculations import decimal_value, money_value
from smr_data import get_smr_read_model
from services.smr_report import _build_report_filename, _author_label


SMR_LAYOUTS = {
    "overview": ("01 · Общий — новый формат", "Люди и все выполненные работы вместе, отдельный раздел каждого объекта."),
    "brigades": ("02 · По бригадам — новый формат", "Отдельный лист каждой бригады: участники и работы по объектам. Общие работы показаны отдельно один раз."),
    "roster": ("03 · Единый состав заявки", "Все участники заявки одним списком вместе с работами. Это оформление файла, а не объединение бригад."),
}
_NAVY = "17324D"
_PALE = "E8F3F4"
_AMBER = "FFF1CB"
_MONTHS = ("января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа", "сентября", "октября", "ноября", "декабря")


def smr_layout_file_entries(app_id: int, objects: list[str]) -> list[dict]:
    return [{
        "kind": "comparison", "variant": key, "team_id": None,
        "team_name": title, "description": description, "objects": objects,
        "download_url": f"/api/kp/apps/{int(app_id)}/smr/download?scope={key}",
    } for key, (title, description) in SMR_LAYOUTS.items()]


def _text(cell, value):
    # Report names/notes may begin with '='. Keep source text as text, not an
    # executable spreadsheet formula (also for imported service names).
    cell.value = str(value or "")
    cell.data_type = "s"


def _band(ws, row, title, *, color=_NAVY):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    cell = ws.cell(row, 1)
    _text(cell, title)
    cell.fill = PatternFill("solid", fgColor=color)
    cell.font = Font(bold=True, color="FFFFFF" if color == _NAVY else _NAVY, size=11)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = max(28, 15 * math.ceil(len(str(title)) / 125))


def _sheet(wb, name):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    for i, width in enumerate((6, 45, 24, 13, 16, 17, 19, 17, 19), 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "D4"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    return ws


def _table_row(ws, row, values, *, header=False):
    height = 27
    for i, value in enumerate(values, 1):
        c = ws.cell(row, i)
        if isinstance(value, (int, float)):
            c.value = value
            c.number_format = "#,##0.00"
        else:
            _text(c, value)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.font = Font(size=11, bold=header, color=_NAVY)
        if header:
            c.fill = PatternFill("solid", fgColor=_PALE)
        width = ws.column_dimensions[get_column_letter(i)].width
        height = max(height, 15 * math.ceil(len(str(value or "")) / max(8, width - 3)))
    ws.row_dimensions[row].height = min(150, height + 5)


def _source_id(row):
    return int(row.get("source_application_id") or row.get("application_id") or 0)


def _team_id(row):
    return int(row.get("team_id") or 0)


def _work_rows(report):
    rows = [
        {**w, "rate_salary": w.get("current_salary", 0), "rate_price": w.get("current_price", 0), "origin": "КП"}
        for w in report.get("plan_works", [])
    ] + [
        {**w, "rate_salary": w.get("salary", 0), "rate_price": w.get("price", 0), "origin": "Доп. работа"}
        for w in report.get("extra_works", [])
    ]
    # The canonical total rounds the sum, not each multiplication. Carry the
    # cent remainder through the rows so displayed line sums match it exactly.
    for rate_key, amount_key in (("rate_salary", "amount_salary"), ("rate_price", "amount_price")):
        cumulative = Decimal("0")
        previous = Decimal("0")
        for row in rows:
            exact = decimal_value(row.get("volume"), field="Объём") * money_value(row.get(rate_key))
            cumulative += exact
            rounded = cumulative.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            amount = rounded - previous
            row[amount_key] = float(amount)
            if amount != exact.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP):
                row["rounding_adjusted"] = True
            previous = rounded
    return rows


def _section(ws, row, context, hours, works, *, financial, participant_salary, roster=False):
    label = context.get("application_label") or context.get("public_number") or f"№{context['id']}"
    _band(ws, row, f"{context.get('object_name') or 'Объект'} · {label}")
    row += 1
    _text(ws.cell(row, 1), "Дата работ")
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=9)
    try:
        date = datetime.strptime(str(context.get("date_target") or "")[:10], "%Y-%m-%d")
        ws.cell(row, 2, date)
        ws.cell(row, 2).number_format = f'dd "{_MONTHS[date.month - 1]}" yyyy'
    except ValueError:
        _text(ws.cell(row, 2), context.get("date_target") or "Не указана")
    row += 2
    _band(ws, row, "Участники и часы", color=_PALE)
    row += 1
    headers = ["№", "Участник", "Исходная бригада" if roster else "Бригада", "Часы"]
    if participant_salary:
        headers.append("ЗП участника, ₽")
    headers.extend(["Вид записи", "Заполнил"])
    _table_row(ws, row, headers, header=True)
    row += 1
    for i, h in enumerate(hours, 1):
        values = [i, h.get("fio") or f"Имя не найдено (карточка {h.get('member_id', '—')})", h.get("team_name") or "Без бригады", float(h.get("hours") or 0)]
        if participant_salary:
            values.append(float(h.get("participant_salary") or 0))
        values.extend(["Допотчёт" if h.get("is_additional") else "Основная", _author_label(h.get("filled_by_fio"), h.get("filled_by_role"))])
        _table_row(ws, row, values)
        row += 1
    if not hours:
        _band(ws, row, "Сохранённых строк часов в этом разделе нет", color=_AMBER)
        row += 1
    _table_row(ws, row, ["", "Человеко-часы", "", sum(float(h.get("hours") or 0) for h in hours)], header=True)
    row += 2
    _band(ws, row, "Все выполненные работы — включая дополнительные", color=_PALE)
    row += 1
    headers = ["№", "Услуга", "Бригада / источник", "Ед. изм.", "Объём"]
    if financial:
        headers.extend(["ЗП за ед., ₽", "Сумма ЗП, ₽", "Цена за ед., ₽", "Стоимость СМР, ₽"])
    _table_row(ws, row, headers, header=True)
    row += 1
    for i, w in enumerate(works, 1):
        shared = not _team_id(w)
        attribution = w.get("team_name") or ("Общая, доли не заданы" if shared else f"Бригада {_team_id(w)}")
        if w.get("is_additional"):
            attribution += " · допотчёт"
        volume = float(w.get("volume") or 0)
        values = [i, w.get("name") or "Без названия", attribution, w.get("unit") or "—", volume]
        if financial:
            salary, price = float(w.get("rate_salary") or 0), float(w.get("rate_price") or 0)
            values.extend([salary, w["amount_salary"], price, w["amount_price"]])
        _table_row(ws, row, values)
        row += 1
    if not works:
        _band(ws, row, "Сохранённых выполненных работ в этом разделе нет.", color=_AMBER)
        row += 1
    return row + 2


async def generate_smr_layout_bytes(db, app_id: int, variant: str, *, include_financial=False, include_participant_salary=False):
    if variant not in SMR_LAYOUTS:
        raise ValueError("Неизвестный вариант СМР")
    report = await get_smr_read_model(db, app_id, include_zero_hours=True)
    contexts = report.get("applications") or []
    if not contexts:
        raise ValueError("Заявка не найдена")
    wb = Workbook()
    wb.remove(wb.active)
    summary = _sheet(wb, "Сводка")
    title, description = SMR_LAYOUTS[variant]
    _band(summary, 1, title)
    _band(summary, 2, description, color=_PALE)
    _band(summary, 3, "Вариант оформления для выбора Кадрами. Не меняет состав, часы, цены и статус СМР в приложении.", color=_AMBER)
    totals = report["totals"]
    metrics = [("Человеко-часы", totals["hours"])]
    if include_participant_salary:
        metrics.append(("ЗП участникам, отдельно от расценок работ", totals["participant_salary"]))
    if include_financial:
        metrics.extend([("ЗП по расценкам работ", totals["salary"]), ("Стоимость СМР", totals["price"])])
    for r, (name, value) in enumerate(metrics, 5):
        _table_row(summary, r, ["", name, value])
    works = _work_rows(report)
    hours = report.get("hours", [])
    row = 6 + len(metrics)
    _band(summary, row, "Объекты и исходные заявки")
    for context in contexts:
        row += 1
        _table_row(summary, row, [context.get("application_label"), context.get("object_name"), context.get("date_target")])
    shared_works = [w for w in works if not _team_id(w)]
    if shared_works:
        _band(summary, row + 2, "Есть общие работы без сохранённых долей. Они учтены один раз; 50/50 автоматически к старым записям не применяется.", color=_AMBER)
    if include_financial and any(w.get("rounding_adjusted") for w in works):
        _band(summary, row + 4, "Округление: остаток копеек распределён между строками, чтобы их сумма точно совпала с итогом СМР.", color=_AMBER)
    sections = []
    if variant == "brigades":
        team_ids = sorted({_team_id(r) for r in hours + works if _team_id(r)})
        for i, team in enumerate(team_ids, 1):
            team_rows = [r for r in hours + works if _team_id(r) == team]
            name = next((r.get("team_name") for r in team_rows if r.get("team_name")), f"Бригада {team}")
            sections.append((f"Бригада {i}", name, [h for h in hours if _team_id(h) == team], [w for w in works if _team_id(w) == team]))
        unassigned_hours = [h for h in hours if not _team_id(h)]
        if shared_works or unassigned_hours:
            sections.append(("Без распределения", "Общие работы / записи без бригады — учтены один раз", unassigned_hours, shared_works))
    else:
        sections.append(("Отчёт", "Единый состав каждой заявки" if variant == "roster" else "Люди и работы по объектам", hours, works))
    for name, section_title, section_hours, section_works in sections:
        ws = _sheet(wb, name)
        _band(ws, 1, section_title)
        row = 3
        for context in contexts:
            ctx_hours = [h for h in section_hours if _source_id(h) == int(context["id"])]
            if variant == "roster":
                ctx_hours.sort(key=lambda h: str(h.get("fio") or "").casefold())
            ctx_works = [w for w in section_works if _source_id(w) == int(context["id"])]
            if variant == "brigades" and not ctx_hours and not ctx_works:
                continue
            row = _section(ws, row, context, ctx_hours, ctx_works, financial=include_financial, participant_salary=include_participant_salary, roster=variant == "roster")
    first = contexts[0]
    object_names = list(dict.fromkeys(c.get("object_name") or "Объект" for c in contexts))
    filename = _build_report_filename(object_name=" + ".join(object_names), app_id=app_id, public_number=first.get("public_number") or first.get("application_label"), date_target=first.get("date_target"), team_name=title.replace(" · ", " "))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), filename
